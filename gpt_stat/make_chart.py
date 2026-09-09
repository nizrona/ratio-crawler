import os
import math
import subprocess
import shutil

from openpyxl import load_workbook

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LABEL_POSITION,
    XL_MARKER_STYLE,
    XL_LEGEND_POSITION,
    XL_TICK_LABEL_POSITION,
)
from pptx.dml.color import RGBColor


# ============================================================
# 1. 파일 설정
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

DATA_FILE = os.path.join(BASE_DIR, "KAIST_수시_경쟁률_입력데이터.xlsx")
OUTPUT_FILE = os.path.join(BASE_DIR, "KAIST_2026_2027_수시지원자_비교.pptx")
OUTPUT_IMAGE_DIR = os.path.join(BASE_DIR, "images")

LAST_YEAR_SHEET = "2026학년도"
THIS_YEAR_SHEET = "2027학년도"


# ============================================================
# 2. 화면 스타일 설정
# ============================================================

FONT_NAME = "맑은 고딕"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

COLOR_LAST_YEAR = RGBColor(120, 120, 120)
COLOR_THIS_YEAR = RGBColor(0, 91, 172)
COLOR_TITLE = RGBColor(30, 30, 30)
COLOR_SUBTITLE = RGBColor(70, 70, 70)

DATA_LABEL_SIZE = 17
AXIS_FONT_SIZE = 12
TITLE_FONT_SIZE = 23
SUBTITLE_FONT_SIZE = 12

MARKER_SIZE = 11

LINE_WIDTH_LAST = Pt(2.2)
LINE_WIDTH_THIS = Pt(3.0)

CATEGORIES = [
    "1일차", "2일차", "3일차",
    "4일차", "5일차", "6일차",
    "7일차", "8일차", "9일차",
]


# ============================================================
# 3. 엑셀 읽기
# ============================================================

def trim_after_first_blank(values):
    result = []
    for value in values:
        if value is None or value == "":
            break
        result.append(int(value))
    return result


def read_year_sheet(ws):
    data = {}
    order = []

    for row in range(2, ws.max_row + 1):
        section = ws.cell(row=row, column=1).value
        if section is None or str(section).strip() == "":
            continue

        section = str(section).strip()
        capacity = ws.cell(row=row, column=2).value

        if capacity is None:
            raise ValueError(
                f"{ws.title} 시트의 '{section}' 모집인원이 비어 있습니다."
            )

        daily_values = [
            ws.cell(row=row, column=col).value
            for col in range(3, 12)
        ]

        values = trim_after_first_blank(daily_values)

        if not values:
            raise ValueError(
                f"{ws.title} 시트의 '{section}' 지원인원이 하나도 없습니다."
            )

        data[section] = {
            "capacity": int(capacity),
            "values": values,
        }
        order.append(section)

    return data, order


def load_input_data(filename):
    wb = load_workbook(filename=filename, data_only=True)

    if LAST_YEAR_SHEET not in wb.sheetnames:
        raise ValueError(f"엑셀에 '{LAST_YEAR_SHEET}' 시트가 없습니다.")

    if THIS_YEAR_SHEET not in wb.sheetnames:
        raise ValueError(f"엑셀에 '{THIS_YEAR_SHEET}' 시트가 없습니다.")

    last_year, last_order = read_year_sheet(wb[LAST_YEAR_SHEET])
    this_year, this_order = read_year_sheet(wb[THIS_YEAR_SHEET])

    missing_this = [s for s in last_order if s not in this_year]
    missing_last = [s for s in this_order if s not in last_year]

    if missing_this:
        raise ValueError("2027학년도 시트에 다음 전형이 없습니다: " + ", ".join(missing_this))

    if missing_last:
        raise ValueError("2026학년도 시트에 다음 전형이 없습니다: " + ", ".join(missing_last))

    return last_year, this_year, this_order


# ============================================================
# 4. 글꼴/텍스트 보조 함수
# ============================================================

def set_run_font(run, size, bold=False, color=None):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_font(paragraph, size, bold=False, color=None):
    for run in paragraph.runs:
        set_run_font(run, size, bold, color)


def add_textbox(slide, text, left, top, width, height, font_size=12, bold=False, color=RGBColor(0, 0, 0)):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_font(run, font_size, bold, color)
    return box


def final_ratio(values, capacity):
    return values[-1] / capacity


# ============================================================
# 5. Y축 라운드 피겨 계산
# ============================================================

def nice_number(value):
    if value <= 0:
        return 1
    exponent = math.floor(math.log10(value))
    fraction = value / (10 ** exponent)

    if fraction <= 1:
        nice_fraction = 1
    elif fraction <= 2:
        nice_fraction = 2
    elif fraction <= 5:
        nice_fraction = 5
    else:
        nice_fraction = 10

    return nice_fraction * (10 ** exponent)


def calculate_axis_scale(values):
    max_value = max(values)
    raw_major = max_value / 6
    major_unit = nice_number(raw_major)

    y_max = math.ceil(max_value / major_unit) * major_unit
    y_max += major_unit
    y_min = -major_unit

    return y_min, y_max, major_unit


# ============================================================
# 6. 차트 시리즈 스타일
# ============================================================

def style_series(series, color, line_width, marker_size):
    series.format.line.color.rgb = color
    series.format.line.width = line_width

    series.marker.style = XL_MARKER_STYLE.CIRCLE
    series.marker.size = marker_size
    series.marker.format.line.color.rgb = color
    series.marker.format.line.width = Pt(2)
    series.marker.format.fill.solid()
    series.marker.format.fill.fore_color.rgb = RGBColor(255, 255, 255)


def style_data_labels(series, font_size, color, position):
    labels = series.data_labels
    labels.show_value = True
    labels.position = position
    labels.number_format = '#,##0'
    labels.font.name = FONT_NAME
    labels.font.size = Pt(font_size)
    labels.font.bold = True
    labels.font.color.rgb = color


def set_point_data_label(point, position, font_size, color):
    dl = point.data_label
    dl.position = position
    dl.font.name = FONT_NAME
    dl.font.size = Pt(font_size)
    dl.font.bold = True
    dl.font.color.rgb = color


# ============================================================
# 7. PPTX → 이미지 변환 함수
# ============================================================

def convert_pptx_to_images(pptx_path, output_folder):
    """
    LibreOffice로 PPTX를 PDF로 변환한 뒤,
    pdf2image로 PDF의 각 페이지를 PNG 이미지로 추출한다.
    """
    os.makedirs(output_folder, exist_ok=True)
    pdf_filename = os.path.splitext(os.path.basename(pptx_path))[0] + ".pdf"
    pdf_path = os.path.join(BASE_DIR, pdf_filename)

    print("\n[변환 1/2] LibreOffice를 이용한 PPTX -> PDF 변환 중...")
    cmd = [
        "libreoffice",
        "--headless",
        "--convert-to",
        "pdf",
        pptx_path,
        "--outdir",
        BASE_DIR,
    ]
    subprocess.run(cmd, check=True)

    print("[변환 2/2] PDF -> 슬라이드별 PNG 이미지 렌더링 중...")
    from pdf2image import convert_from_path

    images = convert_from_path(pdf_path, dpi=200)
    for idx, img in enumerate(images, start=1):
        img_file = os.path.join(output_folder, f"slide_{idx:02d}.png")
        img.save(img_file, "PNG")
        print(f"  - 저장 완료: {img_file}")

    if os.path.exists(pdf_path):
        os.remove(pdf_path)


# ============================================================
# 8. 메인 실행 (데이터 로드 및 PPT 생성)
# ============================================================

last_year, this_year, slide_order = load_input_data(DATA_FILE)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
blank_layout = prs.slide_layouts[6]

for section in slide_order:
    ly = last_year[section]
    ty = this_year[section]
    ly_values = ly["values"]
    ty_values = ty["values"]

    slide = prs.slides.add_slide(blank_layout)

    # 제목
    add_textbox(
        slide,
        f"{section} · 일자별 누적 지원인원 비교",
        0.28, 0.10, 12.4, 0.50,
        TITLE_FONT_SIZE, True, COLOR_TITLE,
    )

    # 상단 요약
    ly_final = ly_values[-1]
    ly_ratio = final_ratio(ly_values, ly["capacity"])
    ty_current = ty_values[-1]
    ty_ratio = final_ratio(ty_values, ty["capacity"])

    summary = (
        f"2026학년도 모집 {ly['capacity']}명 · 최종 {ly_final:,}명 · {ly_ratio:.2f}:1     |     "
        f"2027학년도 모집 {ty['capacity']}명 · 현재 {ty_current:,}명 · {ty_ratio:.2f}:1"
    )

    add_textbox(slide, summary, 0.28, 0.62, 12.3, 0.35, SUBTITLE_FONT_SIZE, False, COLOR_SUBTITLE)

    # 차트 데이터
    chart_data = CategoryChartData()
    chart_data.categories = CATEGORIES

    last_padded = ly_values + [None] * (len(CATEGORIES) - len(ly_values))
    chart_data.add_series("2026학년도", last_padded)

    this_padded = ty_values + [None] * (len(CATEGORIES) - len(ty_values))
    chart_data.add_series("2027학년도", this_padded)

    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(0.55), Inches(1.20), Inches(12.15), Inches(5.80),
        chart_data,
    )

    chart = chart_shape.chart
    chart.has_title = False
    chart.has_legend = True

    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.name = FONT_NAME
    chart.legend.font.size = Pt(12)

    # X축
    category_axis = chart.category_axis
    category_axis.has_title = True
    category_axis.axis_title.text_frame.text = "접수일차"
    for p in category_axis.axis_title.text_frame.paragraphs:
        set_paragraph_font(p, 12, False, COLOR_SUBTITLE)

    category_axis.tick_labels.font.name = FONT_NAME
    category_axis.tick_labels.font.size = Pt(AXIS_FONT_SIZE)
    category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW

    # Y축
    value_axis = chart.value_axis
    all_values = ly_values + ty_values
    y_min, y_max, major_unit = calculate_axis_scale(all_values)

    value_axis.minimum_scale = y_min
    value_axis.maximum_scale = y_max
    value_axis.major_unit = major_unit
    value_axis.tick_labels.number_format = '#,##0;;'
    value_axis.has_title = True
    value_axis.axis_title.text_frame.text = "누적 지원인원"
    for p in value_axis.axis_title.text_frame.paragraphs:
        set_paragraph_font(p, 13, False, COLOR_TITLE)

    value_axis.tick_labels.font.name = FONT_NAME
    value_axis.tick_labels.font.size = Pt(AXIS_FONT_SIZE)
    value_axis.has_major_gridlines = True
    value_axis.has_minor_gridlines = False

    # 시리즈 스타일
    series_last = chart.series[0]
    style_series(series_last, COLOR_LAST_YEAR, LINE_WIDTH_LAST, MARKER_SIZE)
    style_data_labels(series_last, DATA_LABEL_SIZE, COLOR_LAST_YEAR, XL_LABEL_POSITION.ABOVE)

    series_this = chart.series[1]
    style_series(series_this, COLOR_THIS_YEAR, LINE_WIDTH_THIS, MARKER_SIZE)
    style_data_labels(series_this, DATA_LABEL_SIZE, COLOR_THIS_YEAR, XL_LABEL_POSITION.BELOW)

    # 일차별 라벨 배치
    for i in range(len(CATEGORIES)):
        v_last = last_padded[i]
        v_this = this_padded[i]
        pt_last = series_last.points[i]
        pt_this = series_this.points[i]

        if v_this is None:
            set_point_data_label(pt_last, XL_LABEL_POSITION.ABOVE, DATA_LABEL_SIZE, COLOR_LAST_YEAR)
            continue

        if v_last is not None:
            if v_this >= v_last:
                set_point_data_label(pt_this, XL_LABEL_POSITION.ABOVE, DATA_LABEL_SIZE, COLOR_THIS_YEAR)
                set_point_data_label(pt_last, XL_LABEL_POSITION.BELOW, DATA_LABEL_SIZE, COLOR_LAST_YEAR)
            else:
                set_point_data_label(pt_last, XL_LABEL_POSITION.ABOVE, DATA_LABEL_SIZE, COLOR_LAST_YEAR)
                set_point_data_label(pt_this, XL_LABEL_POSITION.BELOW, DATA_LABEL_SIZE, COLOR_THIS_YEAR)


# ============================================================
# 9. PPTX 저장 및 이미지 변환
# ============================================================

prs.save(OUTPUT_FILE)
print("PPTX 생성 완료:", os.path.abspath(OUTPUT_FILE))

# 이미지 추출 실행
convert_pptx_to_images(OUTPUT_FILE, OUTPUT_IMAGE_DIR)
